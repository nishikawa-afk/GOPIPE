"""F-12: 拾い出し → 見積書／材料発注書の生成。

`pricer.Pricer` で各 TakeoffItem に単価を当て、見積明細（材工分離）と
合計（小計＋諸経費＋消費税）を組み立て、見積書・材料発注書の Excel を出力する。
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TakeoffItem
from .pricer import PriceQuote, Pricer

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")
_CENTER = Alignment(horizontal="center", vertical="center")
_YEN = "#,##0"


@dataclass
class EstimateLine:
    """見積 1 行（拾い出し項目 + 単価引き結果）。"""

    item: TakeoffItem
    quote: PriceQuote | None

    @property
    def priced(self) -> bool:
        return self.quote is not None

    @property
    def unit_price(self) -> float:
        return self.quote.unit_price if self.quote else 0.0

    @property
    def amount(self) -> int:
        return int(round(self.quote.total(self.item.quantity))) if self.quote else 0

    @property
    def material(self) -> int:
        return int(round(self.quote.material(self.item.quantity))) if self.quote else 0

    @property
    def labor(self) -> int:
        return int(round(self.quote.labor(self.item.quantity))) if self.quote else 0

    @property
    def man_hours(self) -> float:
        return self.quote.man_hours(self.item.quantity) if self.quote else 0.0


@dataclass
class Estimate:
    """見積全体。小計→諸経費→消費税→合計を派生プロパティで持つ。"""

    lines: list[EstimateLine]
    overhead_rate: float = 0.10  # 諸経費率
    tax_rate: float = 0.10  # 消費税率
    currency_symbol: str = "¥"  # 通貨記号（海外展開: $ 等）

    @property
    def subtotal(self) -> int:
        return int(sum(ln.amount for ln in self.lines))

    @property
    def overhead(self) -> int:
        return int(round(self.subtotal * self.overhead_rate))

    @property
    def site_overhead(self) -> int:
        """現場管理費（諸経費の内数・小計の5%目安、上限は諸経費率）。"""
        return int(round(self.subtotal * min(self.overhead_rate, 0.05)))

    @property
    def general_overhead(self) -> int:
        """一般管理費（諸経費の残り）。"""
        return self.overhead - self.site_overhead

    @property
    def material_total(self) -> int:
        return int(sum(ln.material for ln in self.lines))

    @property
    def labor_total(self) -> int:
        return int(sum(ln.labor for ln in self.lines))

    @property
    def man_hours(self) -> float:
        """総人工（歩掛のあるエントリのみ集計。0なら未計上）。"""
        return round(sum(ln.man_hours for ln in self.lines), 1)

    @property
    def total_ex_tax(self) -> int:
        return self.subtotal + self.overhead

    @property
    def tax(self) -> int:
        return int(round(self.total_ex_tax * self.tax_rate))

    @property
    def total(self) -> int:
        return self.total_ex_tax + self.tax

    @property
    def unpriced(self) -> list[EstimateLine]:
        return [ln for ln in self.lines if not ln.priced]


def build_estimate(
    items: list[TakeoffItem],
    pricer: Pricer,
    *,
    overhead_rate: float = 0.10,
    tax_rate: float = 0.10,
    currency_symbol: str = "¥",
    vendor_id: str | None = None,
) -> Estimate:
    """拾い出し項目に単価を当てて見積を組み立てる。"""
    lines = [EstimateLine(item=it, quote=pricer.quote(it, vendor_id=vendor_id)) for it in items]
    return Estimate(
        lines=lines, overhead_rate=overhead_rate, tax_rate=tax_rate, currency_symbol=currency_symbol,
    )


def _style_header(ws, ncol: int) -> None:
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER


def _yen_column(ws, col_letter: str) -> None:
    for cell in ws[col_letter]:
        cell.number_format = _YEN


def write_estimate_excel(estimate: Estimate, out_path: str | Path) -> Path:
    """見積書 Excel（明細 + 小計/諸経費/消費税/合計）を出力する。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "見積明細"
    header = ["No", "カテゴリ", "名称", "仕様", "場所", "数量", "単位",
              "単価", "金額", "材料費", "労務費", "備考"]
    ws.append(header)
    _style_header(ws, len(header))

    lines = sorted(estimate.lines, key=lambda ln: (ln.item.category or "zzz", ln.item.name))
    for i, ln in enumerate(lines, start=1):
        if not ln.priced:
            note = "単価未設定"
        elif ln.item.confidence >= 0.7:
            note = ""
        else:
            note = f"要確認(信頼度{ln.item.confidence:.2f})"
        ws.append([
            i, ln.item.category or "", ln.item.name, ln.item.spec or "", ln.item.location or "",
            ln.item.quantity, ln.item.unit, ln.unit_price, ln.amount, ln.material, ln.labor, note,
        ])

    # 合計ブロック
    ws.append([])
    summary = [
        ("小計", estimate.subtotal),
        (f"諸経費({estimate.overhead_rate:.0%})", estimate.overhead),
        (f"消費税({estimate.tax_rate:.0%})", estimate.tax),
        ("合計（税込）", estimate.total),
    ]
    for label, val in summary:
        ws.append(["", "", "", "", "", "", "", "", val, "", "", label])
        ws.cell(row=ws.max_row, column=12).font = Font(bold=True)

    for col in ("H", "I", "J", "K"):
        _yen_column(ws, col)
    widths = [5, 12, 22, 18, 14, 8, 6, 11, 13, 12, 12, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    return out_path


def write_purchase_order_excel(estimate: Estimate, out_path: str | Path) -> Path:
    """材料発注書 Excel（材料費ベース明細 + カテゴリ別材料費集計）を出力する。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "材料発注書"
    header = ["No", "カテゴリ", "名称", "仕様", "数量", "単位", "材料費", "備考"]
    ws.append(header)
    _style_header(ws, len(header))

    priced = [ln for ln in estimate.lines if ln.priced and ln.material > 0]
    priced.sort(key=lambda ln: (ln.item.category or "zzz", ln.item.name))
    for i, ln in enumerate(priced, start=1):
        ws.append([
            i, ln.item.category or "", ln.item.name, ln.item.spec or "",
            ln.item.quantity, ln.item.unit, ln.material, (ln.quote.notes if ln.quote else ""),
        ])
    _yen_column(ws, "G")

    # カテゴリ別 材料費集計
    ws2 = wb.create_sheet("カテゴリ別 材料費")
    ws2.append(["カテゴリ", "件数", "材料費合計"])
    _style_header(ws2, 3)
    by_cat: dict[str, list] = defaultdict(lambda: [0, 0])
    for ln in priced:
        cat = ln.item.category or "その他"
        by_cat[cat][0] += 1
        by_cat[cat][1] += ln.material
    for cat, (n, mat) in by_cat.items():
        ws2.append([cat, n, int(mat)])
    _yen_column(ws2, "C")

    for i, w in enumerate([5, 12, 22, 18, 8, 6, 13, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, w in enumerate([14, 8, 16], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    return out_path
