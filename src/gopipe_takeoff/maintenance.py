"""F-17 予防保全サブスク：配管台帳 → 更新時期予測 → 定期点検プラン提案。

拾い出し（施工）データに布設年を与え、管種別の標準耐用年数から更新推奨年・残存年を
算定して配管台帳を作る。フロー型（一発受注）をストック型（定期点検サブスク）へ
転換するための提案書も生成する。
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TakeoffItem

_PATH = Path(__file__).resolve().parents[2] / "prompts" / "service_life.yaml"
_TARGET_CATEGORIES = ("機器", "衛生器具")


def _norm(s: str) -> str:
    return unicodedata.normalize("NFKC", s or "").replace(" ", "").replace("　", "").lower()


@dataclass
class Plan:
    name: str
    interval_months: int
    monthly_fee: int
    scope: str = ""


@dataclass
class ServiceLifeTable:
    rules: list  # list of (keywords: tuple[str,...], years: int, label: str)
    default_years: int = 20

    def lookup(self, item: TakeoffItem) -> tuple[int, str]:
        hay = _norm(f"{item.name} {item.spec or ''}")
        for kws, years, label in self.rules:
            if any(_norm(k) in hay for k in kws):
                return years, label
        return self.default_years, "（標準）"


def load_service_life(path: str | Path | None = None) -> tuple[ServiceLifeTable, list[Plan]]:
    path = Path(path) if path else _PATH
    with open(path, encoding="utf-8") as f:
        data = yaml.safe_load(f) or {}
    rules = []
    for r in data.get("pipe_service_life", []) or []:
        rules.append(
            (tuple(r.get("keywords", []) or []), int(r.get("years", 20)), r.get("label", ""))
        )
    table = ServiceLifeTable(rules=rules, default_years=int(data.get("default_years", 20)))
    plans = [
        Plan(
            name=p.get("name", ""),
            interval_months=int(p.get("interval_months", 12)),
            monthly_fee=int(p.get("monthly_fee", 0)),
            scope=p.get("scope", ""),
        )
        for p in (data.get("plans", []) or [])
    ]
    return table, plans


@dataclass
class LedgerRow:
    item: TakeoffItem
    label: str  # 管種ラベル
    service_life: int  # 耐用年数
    installed_year: int

    @property
    def recommend_year(self) -> int:
        return self.installed_year + self.service_life

    def remaining(self, current_year: int) -> int:
        return self.recommend_year - current_year


def build_ledger(
    items: list[TakeoffItem], installed_year: int, table: ServiceLifeTable
) -> list[LedgerRow]:
    """配管（m単位）と機器を対象に、布設年＋耐用年数の台帳行を作る。"""
    rows: list[LedgerRow] = []
    for it in items:
        if it.unit == "m" or it.category in _TARGET_CATEGORIES:
            years, label = table.lookup(it)
            rows.append(
                LedgerRow(item=it, label=label, service_life=years, installed_year=installed_year)
            )
    return rows


def recommend_plan(ledger: list[LedgerRow], plans: list[Plan], *, current_year: int) -> Plan | None:
    """残存年が短いほど手厚いプランを推奨する。"""
    if not plans:
        return None
    by_intensity = sorted(plans, key=lambda p: p.interval_months)  # 先頭=最も手厚い
    if not ledger:
        return by_intensity[-1]
    min_remaining = min(r.remaining(current_year) for r in ledger)
    if min_remaining <= 0:
        return by_intensity[0]
    if min_remaining <= 5:
        return by_intensity[min(1, len(by_intensity) - 1)]
    return by_intensity[-1]


def _yen(n: int) -> str:
    return f"¥{n:,}"


def write_ledger_excel(ledger: list[LedgerRow], out_path: str | Path, *, current_year: int) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "配管台帳"
    header = ["No", "系統", "名称", "管種", "口径/仕様", "数量", "単位",
              "布設年", "耐用年数", "更新推奨年", "残存年", "状態"]
    ws.append(header)
    hf = Font(bold=True, color="FFFFFF")
    hfill = PatternFill("solid", fgColor="305496")
    ce = Alignment(horizontal="center", vertical="center")
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = hf
        cell.fill = hfill
        cell.alignment = ce
    rows = sorted(ledger, key=lambda r: r.remaining(current_year))
    for i, r in enumerate(rows, start=1):
        rem = r.remaining(current_year)
        state = "★更新時期" if rem <= 0 else ("要注意" if rem <= 5 else "")
        ws.append([
            i, r.item.category or "", r.item.name, r.label, r.item.spec or "",
            r.item.quantity, r.item.unit, r.installed_year, r.service_life,
            r.recommend_year, rem, state,
        ])
    for i, w in enumerate([5, 10, 16, 18, 16, 7, 6, 8, 9, 11, 8, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(out_path)
    return out_path


def build_maintenance_proposal(
    ledger: list[LedgerRow],
    plan: Plan | None,
    plans: list[Plan],
    *,
    installed_year: int,
    current_year: int,
    company: str = "",
) -> str:
    """予防保全プラン提案書（Markdown）を組み立てる。"""
    L: list[str] = []
    L.append(f"# 予防保全プラン提案{('　' + company) if company else ''}")
    L.append("")
    L.append(f"対象物件の配管を台帳化し、更新時期を予測しました"
             f"（布設 {installed_year}年 → 評価 {current_year}年）。")
    L.append("")
    L.append("## 配管台帳（残存年の短い順）")
    L.append("")
    L.append("| 名称 | 管種 | 口径/仕様 | 布設年 | 耐用年数 | 更新推奨年 | 残存年 |")
    L.append("| --- | --- | --- | ---: | ---: | ---: | ---: |")
    rows = sorted(ledger, key=lambda r: r.remaining(current_year))
    for r in rows:
        rem = r.remaining(current_year)
        mark = "★" if rem <= 0 else ""
        L.append(f"| {r.item.name} | {r.label} | {r.item.spec or '—'} | {r.installed_year} "
                 f"| {r.service_life} | {r.recommend_year} | {rem}{mark} |")
    L.append("")
    if rows:
        soonest = rows[0].remaining(current_year)
        if soonest <= 0:
            L.append(f"> ⚠ **更新推奨時期を過ぎた配管があります**（{rows[0].item.name} ほか）。"
                     f"早期の更新計画をおすすめします。")
        else:
            L.append(f"> 最短で **{soonest}年後（{current_year + soonest}年）** に更新推奨の配管があります"
                     f"（{rows[0].item.name}）。")
    L.append("")
    if plan:
        L.append(f"## おすすめプラン: {plan.name}")
        L.append("")
        L.append(f"- 点検周期: **{plan.interval_months}ヶ月ごと**")
        L.append(f"- 料金: **月額 {_yen(plan.monthly_fee)}**（年額 {_yen(plan.monthly_fee * 12)}）")
        L.append(f"- 内容: {plan.scope}")
        L.append("")
    if plans:
        L.append("## 全プラン")
        L.append("")
        L.append("| プラン | 点検周期 | 月額 | 内容 |")
        L.append("| --- | --- | ---: | --- |")
        for p in plans:
            L.append(f"| {p.name} | {p.interval_months}ヶ月 | {_yen(p.monthly_fee)} | {p.scope} |")
        L.append("")
    L.append("## なぜ予防保全か（ストック型への転換）")
    for b in (
        "突発的な漏水・断水・故障を未然に防ぎ、被害と緊急対応コストを削減",
        "更新を計画的に分散でき、まとまった出費を平準化",
        "点検記録が残り、資産価値・売却/賃貸時の説明材料になる",
        "施工業者にとっては「一度きりの工事」を継続的な顧客関係（ストック収益）に変える",
    ):
        L.append(f"- {b}")
    L.append("")
    L.append("---")
    L.append("_本提案は GOPIPE F-17（予防保全サブスク）が拾い出し・布設年から自動生成。"
             "耐用年数は管種の標準値で、実際の更新時期は水質・使用状況で変動します。_")
    return "\n".join(L)


def write_maintenance(
    ledger: list[LedgerRow],
    plan: Plan | None,
    plans: list[Plan],
    out_dir: str | Path,
    *,
    installed_year: int,
    current_year: int,
    company: str = "",
) -> tuple[Path, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    excel = write_ledger_excel(ledger, out_dir / "配管台帳.xlsx", current_year=current_year)
    md = build_maintenance_proposal(
        ledger, plan, plans, installed_year=installed_year, current_year=current_year, company=company
    )
    prop = out_dir / "予防保全プラン提案.md"
    prop.write_text(md, encoding="utf-8")
    return excel, prop
