from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TakeoffItem

HEADER = ["No", "カテゴリ", "名称", "仕様", "場所", "数量", "単位", "ページ", "備考"]


def _note(it: TakeoffItem) -> str:
    """備考列の文言。低信頼は要確認、機器表由来は確定根拠を示す。"""
    if it.confidence < 0.7:
        return f"要確認(信頼度{it.confidence:.2f})"
    if it.source == "reconciled":
        return "機器表で数量確定"
    if it.source == "text_table":
        return "機器表から抽出"
    return ""


def write_excel(items: list[TakeoffItem], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "拾い出し"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    center = Alignment(horizontal="center", vertical="center")

    ws.append(HEADER)
    for col_idx, _ in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # カテゴリ順で並べてから No を振る
    sorted_items = sorted(items, key=lambda x: (x.category or "zzz", x.page, x.name))
    for i, it in enumerate(sorted_items, start=1):
        ws.append(
            [
                i,
                it.category or "",
                it.name,
                it.spec or "",
                it.location or "",
                it.quantity,
                it.unit,
                it.page,
                _note(it),
            ]
        )

    # サマリーシート
    ws2 = wb.create_sheet("カテゴリ別集計")
    ws2.append(["カテゴリ", "件数", "合計数量(同一単位のみ)", "代表単位"])
    by_cat: dict[str, list[TakeoffItem]] = defaultdict(list)
    for it in sorted_items:
        by_cat[it.category or "その他"].append(it)
    for cat, rows in by_cat.items():
        units = {r.unit for r in rows}
        rep_unit = next(iter(units)) if len(units) == 1 else "混在"
        total = sum(r.quantity for r in rows) if rep_unit != "混在" else ""
        ws2.append([cat, len(rows), total, rep_unit])
    for col_idx in range(1, 5):
        c = ws2.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

    widths = [5, 14, 22, 24, 14, 9, 7, 7, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, w in enumerate([16, 8, 18, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    return out_path
